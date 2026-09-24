"""Workbook + SQL contract exercised through the chamber report service."""

from datetime import date, datetime
from types import SimpleNamespace
from zipfile import ZIP_DEFLATED, ZipFile
import re

import openpyxl
import pandas as pd
import pytest
from sqlalchemy import create_engine, text

from src.indicator_domain.application.qtime.chamber_service import ChamberQTimeService
from src.indicator_domain.infrastructure.qtime.chamber_repository import ChamberQTimeRepository
from src.shared_kernel.data_forward import DataForwardPolicy
from src.shared_kernel.report_cutoff import ReportCutoffPolicy
from src.indicator_domain.application.qtime.errors import QTimeDataAccessError
from src.indicator_domain.core.qtime.chamber import summarize_chambers
from src.indicator_domain.application.qtime.chamber_cache import cached_chamber_report, clear_chamber_cache


CHAMBERS = ['PT->OC1', 'OC1->OC2', 'OC2->OC3', 'OC3->OC4', 'OC4->OC5',
            'OC5->OC6', 'OC6->OC7', 'OC7->OC8', 'OC8->MC', 'MC->OC9', 'OC9->OC10']


def write_source(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for col, chamber in enumerate(CHAMBERS, start=5):
        ws.cell(9, col, chamber)
        ws.cell(5, col, '6000')
    ws.cell(9, 2, 'GlassID(P/E)\n基板号(P/E)')
    ws.cell(9, 3, 'LL1Time\nLL1进片时间')
    ws.cell(9, 4, 'LL1->UL1')
    for row_index, (glass, when, duration) in enumerate(rows, start=10):
        ws.cell(row_index, 2, glass)
        ws.cell(row_index, 3, when)
        ws.cell(row_index, 4, 999999)
        for col in range(5, 16):
            ws.cell(row_index, col, duration)
    wb.save(path)
    wb.close()


@pytest.fixture
def report_source(tmp_path):
    engine = create_engine('sqlite+pysqlite:///:memory:')
    with engine.begin() as connection:
        connection.execute(text("ATTACH DATABASE ':memory:' AS eda"))
        connection.execute(text("ATTACH DATABASE ':memory:' AS mdw"))
        connection.execute(text('CREATE TABLE eda.spc_tzbjx_oled (glass_id TEXT, product_spec TEXT, glass_start_time TEXT)'))
        connection.execute(text('CREATE TABLE mdw.dwr_mes_productspec (productspecname TEXT, productcode TEXT)'))
        connection.execute(text("INSERT INTO mdw.dwr_mes_productspec VALUES ('S1','M626'), ('S2','M678'), ('S3','DISABLED')"))
        connection.execute(text("INSERT INTO eda.spc_tzbjx_oled VALUES "
                                "('G1','S1','2026-09-01 08:00:00'), ('G1','S1','2026-09-02 08:00:00'),"
                                "('G2','S2','2026-09-01 08:00:00'), ('G3','S3','2026-09-01 08:00:00'),"
                                "('CONFLICT','S1','2026-09-01 08:00:00'), ('CONFLICT','S3','2026-09-01 08:00:00')"))
    paths = {'3CEE001': tmp_path / 'line1.xlsx', '3CEE002': tmp_path / 'line2.xlsx'}
    yield engine, paths
    engine.dispose()


def make_service(engine, paths, *, forward=0):
    repository = ChamberQTimeRepository(
        SimpleNamespace(engine=engine), paths,
        data_policy=DataForwardPolicy(enabled=bool(forward), offset_days=forward),
        cutoff_policy=ReportCutoffPolicy(),
    )
    return ChamberQTimeService(repository, enabled_products=('M626', 'M678'))


def test_report_matches_glass_products_and_excludes_whole_machine_time(report_source):
    engine, paths = report_source
    when = datetime(2026, 9, 1, 7)
    write_source(paths['3CEE001'], [('G1', when, 6001), ('G1', when, 6001),
                                    ('G3', when, 10), ('UNKNOWN', when, 20), ('CONFLICT', when, 30)])
    write_source(paths['3CEE002'], [('G2', when, 6000)])
    report = make_service(engine, paths).load(as_of=date(2026, 9, 24))
    details = report['details']
    assert len(details) == 22
    assert set(details['prod_code']) == {'M626', 'M678'}
    assert set(details['chamber']) == set(CHAMBERS)
    assert details.query("prod_code == 'M626'")['status'].unique().tolist() == ['超限']
    assert details.query("prod_code == 'M678'")['status'].unique().tolist() == ['正常']
    assert details['duration_seconds'].max() == 6001
    assert report['unmatched_glasses'] == 1
    assert report['ambiguous_glasses'] == 1


def test_missing_measurements_and_repass_do_not_become_zero_or_duplicate(report_source):
    engine, paths = report_source
    write_source(paths['3CEE001'], [('G1', datetime(2026, 9, 1, 7), None),
                                    ('G1', datetime(2026, 9, 2, 7), -1),
                                    ('G1', datetime(2026, 9, 3, 7), 'bad')])
    write_source(paths['3CEE002'], [('G2', datetime(2026, 9, 1, 7), 6001)])
    report = make_service(engine, paths).load(as_of=date(2026, 9, 24))
    details = report['details']
    assert len(details) == 44
    missing = details.query("prod_code == 'M626'")
    assert missing['duration_seconds'].isna().all()
    assert set(missing['status']) == {'缺失'}
    summary = summarize_chambers(details)
    assert summary.query("prod_code == 'M626'")['exceeded_ratio'].isna().all()
    assert summary.query("prod_code == 'M678'")['exceeded_ratio'].eq(1).all()


def test_time_forward_and_today_cutoff_leave_workbook_unchanged(report_source):
    engine, paths = report_source
    write_source(paths['3CEE001'], [('G1', datetime(2026, 9, 20, 12), 1),
                                    ('G1', datetime(2026, 9, 20, 12, 0, 1), 2),
                                    ('G1', datetime(2026, 9, 19, 23), 3)])
    # Serial date representation used by enterprise values-only normalization.
    write_source(paths['3CEE002'], [('G2', 46266.5, 5)])
    before = paths['3CEE001'].read_bytes()
    details = make_service(engine, paths, forward=4).load(as_of=date(2026, 9, 24))['details']
    assert len(details) == 33
    assert details['entry_time'].max() == pd.Timestamp('2026-09-24 12:00:00')
    assert 2 not in details['duration_seconds'].values
    assert details.query("prod_code == 'M678'")['entry_time'].iloc[0] == pd.Timestamp('2026-09-05 12:00:00')
    assert paths['3CEE001'].read_bytes() == before


@pytest.mark.parametrize('failure', ['missing_file', 'duplicate_conflict', 'header', 'target', 'database'])
def test_source_failure_is_not_published_as_an_empty_success(report_source, failure):
    engine, paths = report_source
    rows = [('G1', datetime(2026, 9, 1, 7), 1)]
    if failure == 'duplicate_conflict':
        rows.append(('G1', datetime(2026, 9, 1, 7), 2))
    write_source(paths['3CEE001'], rows)
    write_source(paths['3CEE002'], rows)
    if failure == 'missing_file':
        paths['3CEE002'] = paths['3CEE002'].with_name('absent.xlsx')
    if failure in {'header', 'target'}:
        wb = openpyxl.load_workbook(paths['3CEE001'])
        wb.active.cell(9 if failure == 'header' else 5, 5, 'bad')
        wb.save(paths['3CEE001'])
        wb.close()
    if failure == 'database':
        engine = None
    with pytest.raises(QTimeDataAccessError, match='蒸镀单腔停留时间数据读取失败'):
        make_service(engine, paths).load(as_of=date(2026, 9, 24))


def test_supplier_incorrect_sheet_dimensions_do_not_hide_rows(report_source):
    engine, paths = report_source
    for path in paths.values():
        write_source(path, [('G1', datetime(2026, 9, 1, 7), 12)])
        with ZipFile(path) as archive:
            members = {name: archive.read(name) for name in archive.namelist()}
        members['xl/worksheets/sheet1.xml'] = re.sub(
            rb'<dimension ref="[^"]+"', b'<dimension ref="A1"', members['xl/worksheets/sheet1.xml'],
        )
        with ZipFile(path, 'w', ZIP_DEFLATED) as archive:
            for name, content in members.items():
                archive.writestr(name, content)
    details = make_service(engine, paths).load(as_of=date(2026, 9, 24))['details']
    assert len(details) == 22


def test_each_physical_column_keeps_its_own_chamber_and_target(report_source):
    engine, paths = report_source
    for path in paths.values():
        write_source(path, [('G1', datetime(2026, 9, 1, 7), 1)])
        wb = openpyxl.load_workbook(path)
        for col, value, target in zip(range(5, 16),
                                      [10, 20, 30, 40, 50, 60, 70, 80, 90, 100, 110],
                                      [9, 20, 31, 40, 51, 60, 71, 80, 91, 100, 109]):
            wb.active.cell(10, col, value)
            wb.active.cell(5, col, target)
        wb.save(path)
        wb.close()
    details = make_service(engine, paths).load(as_of=date(2026, 9, 24))['details']
    first = details.query("line == '3CEE001'").set_index('chamber').reindex(CHAMBERS)
    assert first['duration_seconds'].tolist() == [10, 20, 30, 40, 50, 60, 70, 80, 90, 100, 110]
    assert first['target_seconds'].tolist() == [9, 20, 31, 40, 51, 60, 71, 80, 91, 100, 109]
    assert first.loc[first.status.eq('超限')].index.tolist() == ['PT->OC1', 'OC9->OC10']


def test_product_lookup_uses_source_month_across_display_month_boundary(report_source):
    engine, paths = report_source
    for path in paths.values():
        write_source(path, [('G1', datetime(2026, 9, 30, 11), 5)])
    details = make_service(engine, paths, forward=4).load(as_of=date(2026, 10, 4))['details']
    assert len(details) == 22
    assert details.entry_time.unique().tolist() == [pd.Timestamp('2026-10-04 11:00:00')]


def test_file_and_enabled_scope_changes_invalidate_cached_payload(report_source):
    clear_chamber_cache()
    engine, paths = report_source
    for path in paths.values():
        write_source(path, [('G1', datetime(2026, 9, 1, 7), 10), ('G2', datetime(2026, 9, 1, 7), 20)])
    service = make_service(engine, paths)
    first = cached_chamber_report(service, as_of=date(2026, 9, 24))
    assert first['details'].duration_seconds.max() == 20
    write_source(paths['3CEE001'], [('G1', datetime(2026, 9, 1, 7), 9000)])
    changed = cached_chamber_report(service, as_of=date(2026, 9, 24))
    assert changed['details'].duration_seconds.max() == 9000
    assert first['details'].duration_seconds.max() == 20
    service.enabled_products = ('M678',)
    restricted = cached_chamber_report(service, as_of=date(2026, 9, 24))
    assert set(restricted['details'].prod_code) == {'M678'}
    clear_chamber_cache()


def test_explicit_sources_with_equal_business_signatures_do_not_share_cache():
    clear_chamber_cache()

    class Source:
        def __init__(self, duration):
            self.duration = duration

        def cache_signature(self):
            return ('same-business-signature',)

        def read(self, *, as_of):
            return {
                'measurements': pd.DataFrame([{'line': '3CEE001', 'glass_id': 'G',
                    'entry_time': pd.Timestamp('2026-09-01'), **dict.fromkeys(CHAMBERS, self.duration)}]),
                'products': pd.DataFrame({'glass_id': ['G'], 'prod_code': ['M626']}),
                'targets': pd.DataFrame({'line': '3CEE001', 'chamber': CHAMBERS, 'target_seconds': 6000}),
                'invalid_rows': 0,
            }

    first = ChamberQTimeService(Source(1), enabled_products=('M626',))
    second = ChamberQTimeService(Source(2), enabled_products=('M626',))
    assert cached_chamber_report(first, as_of=date(2026, 9, 24))['details'].duration_seconds.unique().tolist() == [1]
    assert cached_chamber_report(second, as_of=date(2026, 9, 24))['details'].duration_seconds.unique().tolist() == [2]
    clear_chamber_cache()
