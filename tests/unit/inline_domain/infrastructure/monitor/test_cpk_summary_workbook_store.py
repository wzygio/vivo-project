from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor

import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.styles import Font

from src.inline_domain.infrastructure.monitor.cpk_summary_workbook_store import (
    CpkSummaryWorkbookStore,
)
from src.inline_domain.infrastructure.monitor.summary_workbook_store import (
    MonitorSummaryWorkbookStore,
)
from src.inline_domain.infrastructure.monitor import summary_workbook_store as store_module
from src.shared_kernel.utils.excel_tools import WorkbookWriteResult
from tests.unit.inline_domain.infrastructure.monitor.test_summary_workbook_store import (
    _current_rows,
    _write_legacy_workbook,
)


def _cpk_row(time_label="2026-Q3", display_label="Q3") -> pd.DataFrame:
    return pd.DataFrame([{
        "产品": "M626", "周期类型": "季度", "时间标签": time_label,
        "显示标签": display_label, "CPK总项目数": 7, "Cpk≥1.33达标率": 0.57,
    }])


@pytest.mark.parametrize("cpk_first", [False, True])
def test_shared_workbook_updates_preserve_closed_rows_and_other_sheet(
    tmp_path, cpk_first
) -> None:
    path = tmp_path / "summary.xlsx"
    _write_legacy_workbook(path)
    with pd.ExcelWriter(path, mode="a", engine="openpyxl", if_sheet_exists="replace") as writer:
        _cpk_row("2026-Q2", "Q2").to_excel(writer, sheet_name="CPK", index=False)
    workbook = load_workbook(path)
    notes = workbook.create_sheet("notes")
    notes["A1"] = "=1+2"
    notes["A1"].font = Font(bold=True)
    workbook.save(path)
    workbook.close()
    alarm = MonitorSummaryWorkbookStore(path)
    cpk = CpkSummaryWorkbookStore(path)
    original = cpk.read().iloc[0]
    assert original["监控类型"] == "SPC"
    assert original["厂别"] == "ALL"
    assert pd.isna(original["达标项目数"])
    assert alarm._update_lock is cpk._update_lock
    updates = [(alarm, _current_rows()), (cpk, _cpk_row())]
    for store, rows in reversed(updates) if cpk_first else updates:
        store.upsert_current_periods(rows, as_of=pd.Timestamp("2026-09-07"))

    result = cpk.read()
    pd.testing.assert_series_equal(
        result.loc[result["时间标签"].eq("2026-Q2")].iloc[0], original,
    )
    assert set(result["时间标签"]) == {"2026-Q2", "2026-Q3"}
    assert alarm.read().loc[lambda frame: frame["时间标签"].eq("2026-Q3"), "过货量"].item() == 40
    workbook = load_workbook(path)
    assert workbook["notes"]["A1"].value == "=1+2"
    assert workbook["notes"]["A1"].font.bold
    workbook.close()


def test_cpk_closed_period_cannot_be_overwritten(tmp_path) -> None:
    store = CpkSummaryWorkbookStore(tmp_path / "summary.xlsx")
    with pytest.raises(ValueError, match="current year/quarter/month/week"):
        store.upsert_current_periods(
            _cpk_row("2026-Q2", "Q2"), as_of=pd.Timestamp("2026-09-07")
        )


def test_simultaneous_alarm_and_cpk_updates_share_workbook_transaction(tmp_path) -> None:
    path = tmp_path / "summary.xlsx"
    _write_legacy_workbook(path)
    with pd.ExcelWriter(path, mode="a", engine="openpyxl", if_sheet_exists="replace") as writer:
        _cpk_row("2026-Q2", "Q2").to_excel(writer, sheet_name="CPK", index=False)
    alarm = MonitorSummaryWorkbookStore(path)
    cpk = CpkSummaryWorkbookStore(path)
    with ThreadPoolExecutor(max_workers=2) as executor:
        futures = [
            executor.submit(store.upsert_current_periods, rows, as_of=pd.Timestamp("2026-09-07"))
            for store, rows in [(alarm, _current_rows()), (cpk, _cpk_row())]
        ]
        for future in futures:
            future.result()
    assert set(cpk.read()["时间标签"]) == {"2026-Q2", "2026-Q3"}
    assert set(alarm.read()["时间标签"]) == {
        "2026", "2026-Q2", "2026-Q3", "2026-08", "2026-09", "2026-W37",
    }


def test_protected_cpk_write_uses_shared_com_writer_and_cpk_contract(
    tmp_path, monkeypatch
) -> None:
    path = tmp_path / "protected.xlsx"
    path.write_bytes(b"protected")
    store = CpkSummaryWorkbookStore(path)
    calls = []

    def write(candidate, frame, *, sheet_name, normalizer):
        calls.append((candidate, sheet_name, normalizer(frame)))
        return WorkbookWriteResult(True, candidate, (sheet_name,))

    monkeypatch.setattr(store_module, "_replace_encrypted_summary_sheet_via_com", write)
    result = store._write(_cpk_row())

    assert result.written
    assert calls[0][:2] == (path, "CPK")
    assert calls[0][2]["CPK总项目数"].item() == 7
    assert pd.isna(calls[0][2]["达标项目数"].item())
