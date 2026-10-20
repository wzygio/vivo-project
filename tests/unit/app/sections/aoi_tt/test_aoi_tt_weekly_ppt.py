"""Verify TT data, labels and eight product columns in the downloadable PPT."""
from datetime import date
from io import BytesIO

import pandas as pd
from pptx import Presentation

from app.sections.inline_domain.aoi_tt.aoi_tt_weekly_export import build_aoi_tt_weekly_ppt
from src.inline_domain.application.aoi_tt.aoi_tt_service import AoiTtReportViewModel
from tests.unit.app.sections.shared.weekly_ppt_fixtures import template


def test_tt_ppt_contains_eight_products_correct_values_and_no_rs_labels():
    products = tuple(f"M{i}" for i in range(8))

    def load(product):
        row = dict(prod_code=product, factory="ARRAY", step_id="12620", tt_name="TDSUM",
                   tt_qty=12, sheet_id="S1", start_time=pd.Timestamp("2026-10-14"), particle_size="Total")
        rows = [] if product == "M7" else [row]
        specs = pd.DataFrame([dict(step_id="12620", tt_name="TDSUM", prod_code=product)])
        return AoiTtReportViewModel(pd.DataFrame(rows), specs, pd.DataFrame())

    content = build_aoi_tt_weekly_ppt(products, factories=("ARRAY",), steps=(),
        load_report=load, reference_date=date(2026, 10, 20), template_bytes=template())
    deck = Presentation(BytesIO(content))
    assert len(deck.slides) == 1
    grid = next(s.table for s in deck.slides[0].shapes if s.has_table)
    assert [cell.text for cell in grid.rows[0].cells] == ["分类", "层别", "项目", *products]
    assert [cell.text for cell in grid.rows[1].cells] == ["TT Defect", "M1", "TDSUM（Total）", *["12"] * 7, "无过货"]
    assert deck.core_properties.title == "AOI_TT 周度报告"
    assert deck.core_properties.subject == "平均每片 TT 个数"
    text = " ".join(s.text for s in deck.slides[0].shapes if s.has_text_frame)
    assert "平均每片 TT 个数" in text and "AOI_TT" in text
    assert "RS" not in text


def test_tt_empty_export_has_tt_business_message():
    empty = AoiTtReportViewModel(pd.DataFrame(), pd.DataFrame(), pd.DataFrame())
    content = build_aoi_tt_weekly_ppt(("M626",), factories=("OLED",), steps=(),
        load_report=lambda _: empty, reference_date=date(2026, 10, 20), template_bytes=template())
    deck = Presentation(BytesIO(content))
    text = " ".join(s.text for s in deck.slides[0].shapes if s.has_text_frame)
    assert "暂无 TT Defect 项目" in text and "RS" not in text
