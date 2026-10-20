"""Weekly PPTs preserve template assets and product/Code cell placement."""
from datetime import date
from io import BytesIO
from types import SimpleNamespace
import base64
import pytest

from pptx import Presentation
from pptx.util import Inches

from app.charts.inline_domain.aoi_rs_weekly_ppt import build_weekly_rs_ppt
from tests.unit.app.sections.shared.weekly_ppt_fixtures import template




def table(*, factory="ARRAY", count=1, products=("M626", "M678")):
    rows = tuple(SimpleNamespace(
        step_id="12629", layer="M1", code=f"A{i:04}", label=f"A{i:04}-测试缺陷",
        values=("0.125", "无过货")[:len(products)],
    ) for i in range(count))
    return SimpleNamespace(factory=factory, products=products, rows=rows,
                           week_start=date(2026, 10, 12), week_end=date(2026, 10, 19))


def test_ppt_preserves_template_and_numeric_cells():
    source = template()
    result = build_weekly_rs_ppt([table()], template_bytes=source)
    deck = Presentation(BytesIO(result))
    assert len(deck.slides) == 1
    slide = deck.slides[0]
    assert "Template brand" in [s.text for s in slide.shapes if s.has_text_frame]
    grid = next(s.table for s in slide.shapes if s.has_table)
    assert [c.text for c in grid.rows[0].cells] == ["分类", "层别", "项目", "M626", "M678"]
    assert [c.text for c in grid.rows[1].cells] == ["RS Code", "M1", "A0000-测试缺陷", "0.125", "无过货"]
    assert "2026-10-12" in " ".join(s.text for s in slide.shapes if s.has_text_frame)


def test_pagination_repeats_template_and_headers_without_losing_codes():
    result = build_weekly_rs_ppt([table(count=23)], template_bytes=template())
    deck = Presentation(BytesIO(result))
    assert len(deck.slides) >= 3
    labels = []
    for slide in deck.slides:
        assert "Template brand" in [s.text for s in slide.shapes if s.has_text_frame]
        grid = next(s.table for s in slide.shapes if s.has_table)
        assert grid.cell(0, 1).text == "层别"
        labels.extend(grid.cell(r, 2).text for r in range(1, len(grid.rows)))
        shape = next(s for s in slide.shapes if s.has_table)
        assert shape.top + shape.height < deck.slide_height
    assert labels == [f"A{i:04}-测试缺陷" for i in range(23)]


def test_non_array_omits_layer_and_empty_factory_has_safe_empty_state():
    result = build_weekly_rs_ppt([table(factory="OLED"), table(factory="TP", count=0)],
                               template_bytes=template())
    deck = Presentation(BytesIO(result))
    grid = next(s.table for s in deck.slides[0].shapes if s.has_table)
    assert [c.text for c in grid.rows[0].cells] == ["分类", "项目", "M626", "M678"]
    text = " ".join(s.text for s in deck.slides[1].shapes if s.has_text_frame)
    assert "暂无" in text


@pytest.mark.parametrize("product_count, expected_pages", [(8, 1), (9, 2)])
@pytest.mark.parametrize("factory, first_product_column", [("ARRAY", 3), ("OLED", 2)])
def test_eight_products_share_page_and_overflow_preserves_images_and_cells(
    product_count, expected_pages, factory, first_product_column,
):
    source = Presentation(BytesIO(template()))
    png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Wl2zGQAAAAASUVORK5CYII=")
    source.slides[0].shapes.add_picture(BytesIO(png), Inches(10), Inches(.1), Inches(.3), Inches(.3))
    stream = BytesIO()
    source.save(stream)
    products = tuple(f"M{i}" for i in range(product_count))
    sample = table(products=products, factory=factory)
    sample.rows[0].values = tuple(str(i) for i in range(product_count))
    result = build_weekly_rs_ppt([sample], template_bytes=stream.getvalue())
    deck = Presentation(BytesIO(result))
    assert len(deck.slides) == expected_pages
    seen = []
    for slide in deck.slides:
        picture = next(shape for shape in slide.shapes if shape.shape_type == 13)
        assert picture.image.blob == png
        grid = next(shape.table for shape in slide.shapes if shape.has_table)
        assert len(grid.columns) - first_product_column <= 8
        seen.extend((grid.cell(0, c).text, grid.cell(1, c).text)
                    for c in range(first_product_column, len(grid.columns)))
    assert seen == [(f"M{i}", str(i)) for i in range(product_count)]
