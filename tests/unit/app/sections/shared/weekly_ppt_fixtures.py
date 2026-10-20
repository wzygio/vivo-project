"""In-memory Open XML fixture requiring neither Office nor installed PPT files."""
from io import BytesIO
from zipfile import ZipFile
from pptx import Presentation
from pptx.util import Inches


def template():
    # Build Open XML in memory: corporate file hooks can encrypt installed .pptx
    # assets, including python-pptx's default template. Tests require no Office.
    ns = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
    tree = '<p:spTree><p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr/></p:spTree>'
    relns = 'http://schemas.openxmlformats.org/package/2006/relationships'
    reltype = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
    def rels(items):
        return f'<Relationships xmlns="{relns}">' + ''.join(f'<Relationship Id="{rid}" Type="{reltype}{kind}" Target="{target}"/>' for rid, kind, target in items) + '</Relationships>'
    types = {'/ppt/presentation.xml': 'presentation.main', '/ppt/slideMasters/slideMaster1.xml': 'slideMaster', '/ppt/slideLayouts/slideLayout1.xml': 'slideLayout'}
    files = {
        '[Content_Types].xml': '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>' + ''.join(f'<Override PartName="{part}" ContentType="application/vnd.openxmlformats-officedocument.presentationml.{kind}+xml"/>' for part, kind in types.items()) + '</Types>',
        '_rels/.rels': rels([('rId1', 'officeDocument', 'ppt/presentation.xml')]),
        'ppt/presentation.xml': f'<p:presentation {ns}><p:sldMasterIdLst><p:sldMasterId id="2147483648" r:id="rId1"/></p:sldMasterIdLst><p:sldIdLst/><p:sldSz cx="12192000" cy="6858000"/><p:notesSz cx="6858000" cy="9144000"/></p:presentation>',
        'ppt/_rels/presentation.xml.rels': rels([('rId1', 'slideMaster', 'slideMasters/slideMaster1.xml')]),
        'ppt/slideMasters/slideMaster1.xml': f'<p:sldMaster {ns}><p:cSld>{tree}</p:cSld><p:clrMap accent1="accent1" accent2="accent2" accent3="accent3" accent4="accent4" accent5="accent5" accent6="accent6" bg1="lt1" bg2="lt2" folHlink="folHlink" hlink="hlink" tx1="dk1" tx2="dk2"/><p:sldLayoutIdLst><p:sldLayoutId id="1" r:id="rId1"/></p:sldLayoutIdLst><p:txStyles/></p:sldMaster>',
        'ppt/slideMasters/_rels/slideMaster1.xml.rels': rels([('rId1', 'slideLayout', '../slideLayouts/slideLayout1.xml')]),
        'ppt/slideLayouts/slideLayout1.xml': f'<p:sldLayout {ns} type="blank"><p:cSld name="Blank">{tree}</p:cSld><p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr></p:sldLayout>',
        'ppt/slideLayouts/_rels/slideLayout1.xml.rels': rels([('rId1', 'slideMaster', '../slideMasters/slideMaster1.xml')]),
    }
    stream = BytesIO()
    with ZipFile(stream, 'w') as package:
        for name, content in files.items():
            package.writestr(name, content)
    deck = Presentation(BytesIO(stream.getvalue()))
    deck.slide_width, deck.slide_height = Inches(13.3333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[0])
    slide.shapes.add_textbox(Inches(11), 0, Inches(2), Inches(.5)).text = "Template brand"
    output = BytesIO()
    deck.save(output)
    return output.getvalue()

