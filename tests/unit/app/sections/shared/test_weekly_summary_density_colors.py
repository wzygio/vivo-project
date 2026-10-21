"""Exported font colors use each product's unrounded density and matching spec."""
from datetime import date
from io import BytesIO

import pandas as pd
import pytest
from pptx import Presentation

from app.sections.inline_domain.aoi_rs.aoi_rs_weekly_export import build_aoi_rs_weekly_ppt
from app.sections.inline_domain.aoi_tt.aoi_tt_weekly_export import build_aoi_tt_weekly_ppt
from src.inline_domain.application.aoi_rs.aoi_rs_service import AoiRsReportViewModel
from src.inline_domain.application.aoi_tt.aoi_tt_service import AoiTtReportViewModel
from src.shared_kernel.config import ConfigLoader
from tests.unit.app.sections.shared.weekly_ppt_fixtures import template

GREEN, RED, BLACK = "008000", "FF0000", "000000"


def product_cells(content, offset):
    deck = Presentation(BytesIO(content))
    result = []
    for slide in deck.slides:
        grid = next(s.table for s in slide.shapes if s.has_table)
        result.append([(grid.cell(1, col).text,
                        str(grid.cell(1, col).text_frame.paragraphs[0].font.color.rgb))
                       for col in range(offset, len(grid.columns))])
    return result


def test_rs_colors_match_week_spec_and_product_across_eight_column_boundary(monkeypatch):
    monkeypatch.setattr(ConfigLoader, "get_aoi_rs_special_decoration_factories", lambda: ())
    products = tuple(f"M{i}" for i in range(9))
    quantities = (1, 1, 1, 0, 1, 1.00000001, 1, 1, 1)
    limits = (1, .5, None, 0, 1, 1, 2, .5, .25)

    def load(product):
        i = products.index(product)
        fact = dict(prod_code=product, factory="ARRAY", step_id="12629", rs_code="R1",
                    code_qty=quantities[i], start_time=pd.Timestamp("2026-10-14"), sheet_id="S1")
        details = pd.DataFrame([] if i == 4 else [fact])
        passes = pd.DataFrame([] if i == 4 else [fact])
        specs = [dict(prod_code=product, factory="ARRAY", step_id="12629", rs_code="R1",
                      type_flag="MWD_RATIO", spec=limits[i]),
                 dict(prod_code=product, factory="ARRAY", step_id="12629", rs_code="R1",
                      type_flag="LOT_RATIO", spec=1000)]
        return AoiRsReportViewModel(details, passes, pd.DataFrame(specs), pd.DataFrame(),
                                    pd.DataFrame(), pd.DataFrame())

    content = build_aoi_rs_weekly_ppt(products, factories=("ARRAY",), steps=(), load_report=load,
        reference_date=date(2026, 10, 21), template_bytes=template())
    assert product_cells(content, 3) == [
        [("1", GREEN), ("1", RED), ("1", BLACK), ("0", GREEN),
         ("无过货", BLACK), ("1", RED), ("1", GREEN), ("1", RED)],
        [("1", RED)],
    ]


@pytest.mark.parametrize("size", ["Total", "S"])
def test_tt_colors_use_usl_independently_of_ucl_and_keep_missing_data_black(size):
    products = tuple(f"M{i}" for i in range(6))

    def load(product):
        i = products.index(product)
        base = dict(prod_code=product, factory="ARRAY", step_id="12620", tt_name="TDSUM",
                    tt_qty=8, start_time=pd.Timestamp("2026-10-14"), sheet_id="S1", particle_size=size)
        rows = [] if i == 3 else [base]
        if i == 4:
            rows = [{**base, "tt_name": "OTHER", "particle_size": "Total"}]
        specs = pd.DataFrame([dict(prod_code=product, step_id="12620", tt_name="TDSUM",
                                   usl=[8, 7, None, 8, 8, 9][i], ucl=1)])
        return AoiTtReportViewModel(pd.DataFrame(rows), specs, pd.DataFrame())

    content = build_aoi_tt_weekly_ppt(products, factories=("ARRAY",), steps=(), particle_sizes=(size,),
        load_report=load, reference_date=date(2026, 10, 21), template_bytes=template())
    # TDSUM sorts after OTHER, so select its row explicitly.
    deck = Presentation(BytesIO(content))
    grid = next(s.table for s in deck.slides[0].shapes if s.has_table)
    row_index = next(i for i in range(1, len(grid.rows)) if "TDSUM" in grid.cell(i, 2).text)
    actual = [(grid.cell(row_index, c).text,
               str(grid.cell(row_index, c).text_frame.paragraphs[0].font.color.rgb))
              for c in range(3, len(grid.columns))]
    assert actual == [("8", GREEN), ("8", RED), ("8", BLACK),
                      ("无过货", BLACK), ("数据不足", BLACK), ("8", GREEN)]
