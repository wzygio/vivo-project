from datetime import datetime
from io import BytesIO

import openpyxl
import pandas as pd
import pytest

from src.yield_domain.infrastructure.period_trend_source import PeriodTrendDatabaseSource, BASE
from src.yield_domain.infrastructure.period_trend_workbook import PeriodTrendWorkbook
from src.yield_domain.core.mwd_trend.period_trend import RATE_COLUMNS
from app.sections.yield_domain.period_trend_data import export_period_tables


def source(tmp_path):
    path = tmp_path / 'modifier.xlsx'
    path.write_bytes(b'test')
    obj = PeriodTrendDatabaseSource(path)
    obj.catalog = pd.DataFrame([['G', d, d] for d in ['A', 'B', 'C']], columns=['defect_group', 'defect_desc', 'defect_code'])
    obj._frames_signature = obj.maintenance_signature
    obj._frames = {'P_Code级': pd.DataFrame([
        ['A', '月度', '2026-06', .1, .2],
        ['B', '月度', '2026-06', None, None],
        ['C', '月度', '2026-05', .1, .3],
    ], columns=['不良类型', '周期类型', '时间标签', '当月良损', '指定良损']),
    'P_Group级': pd.DataFrame([['G', '月度', '2026-06', .4, None]], columns=['不良类型', '周期类型', '时间标签', '当月良损', '指定良损'])}
    return obj


def test_maintained_missing_is_nan_and_group_uses_own_rate(tmp_path):
    result = source(tmp_path).maintained_rates('P', '2026-06').set_index('defect_code')
    assert result.loc['A', 'defect_rate'] == .2
    assert pd.isna(result.loc['B', 'defect_rate'])
    assert result.loc['C', 'defect_rate'] == .3
    assert result.loc['', 'defect_rate'] == .4


def test_source_signature_cannot_publish_stale_frames(tmp_path, monkeypatch):
    obj = source(tmp_path)
    new = {'changed': pd.DataFrame()}
    monkeypatch.setattr('src.yield_domain.infrastructure.period_trend_source.read_workbook_sheets', lambda p: new)
    obj.path.write_bytes(b'updated')
    import os
    stat = obj.path.stat()
    os.utime(obj.path, ns=(stat.st_atime_ns, stat.st_mtime_ns + 1000000))
    assert obj.frames is new


def test_bound_month_query_has_no_defect_join_and_keeps_microseconds(tmp_path, monkeypatch):
    obj = source(tmp_path)
    monkeypatch.setattr('src.yield_domain.infrastructure.period_trend_source.ConfigLoader.get_enabled_products', lambda: ['P'])
    from types import SimpleNamespace
    monkeypatch.setattr('src.yield_domain.infrastructure.period_trend_source.ConfigLoader.load_config', lambda p: SimpleNamespace(data_source=SimpleNamespace(work_order_types=["P' OR 1=1 --"])))
    seen = []
    def read(sql, params, **kwargs):
        seen.append((sql, params))
        return pd.DataFrame({'total_panels': [10]})
    obj._read = read
    assert obj.throughput('P', '2026-10', datetime(2026, 10, 22)) == 10
    sql, params = seen[0]
    assert 'dws_dft' not in sql and "OR 1=1" not in sql
    assert params['start'] == '20260927'
    assert params['end'] == '20261018'
    assert params['cutoff'] == '20261018120000000000'
    assert "rpad(dwp.first_ship_time,20,'0')" in sql


def test_atomic_history_cas_and_failed_replace_preserve_file(tmp_path, monkeypatch):
    store = PeriodTrendWorkbook(tmp_path / 'history.xlsx')
    rates = pd.DataFrame([['2026-01', 'Group', 'G', '', '', .1]], columns=RATE_COLUMNS)
    inputs = pd.DataFrame({'month': ['2026-01'], 'total_panels': [100]})
    store.save('A', rates, inputs, {})
    before = store.path.read_bytes()
    with pytest.raises(RuntimeError, match='concurrently'):
        store.save('A', rates.assign(defect_rate=.2), inputs, {})
    assert store.path.read_bytes() == before
    _, _, meta = store.read('A')
    meta['expected_version'] = meta['version']
    monkeypatch.setattr('src.yield_domain.infrastructure.period_trend_workbook.os.replace', lambda *args: (_ for _ in ()).throw(OSError('disk failed')))
    with pytest.raises(OSError):
        store.save('A', rates.assign(defect_rate=.2), inputs, meta)
    assert store.path.read_bytes() == before
    assert not list(tmp_path.glob('.period-*'))


def test_export_is_public_and_formula_safe():
    table = pd.DataFrame([{'指标': 'Code', 'Defect Group': 'G', 'Defect Code': '=1+1', '缺陷名称': 'test', '2026-H1': .1, 'flag': 'INTERNAL'}])
    data = export_period_tables({'A': table, 'B': table})
    book = openpyxl.load_workbook(BytesIO(data))
    assert book.sheetnames == ['A', 'B']
    assert book['A']['C2'].data_type == 's'
    assert book['A']['C2'].value == '=1+1'
    assert 'INTERNAL' not in str(list(book['A'].values))


@pytest.mark.parametrize('value', [200, 'abc', -1, float('inf')])
def test_invalid_maintained_rate_is_rejected(tmp_path, value):
    obj = source(tmp_path)
    obj._frames['P_Code级']['指定良损'] = obj._frames['P_Code级']['指定良损'].astype(object)
    obj._frames['P_Code级'].loc[0, '指定良损'] = value
    with pytest.raises(ValueError):
        obj.maintained_rates('P', '2026-06')


def _concurrent_publish(path, product, expected, wait, signal, queue):
    store = PeriodTrendWorkbook(path)
    rates = pd.DataFrame([['2026-01', 'Group', 'G', '', '', .2]], columns=RATE_COLUMNS)
    inputs = pd.DataFrame({'month': ['2026-01'], 'total_panels': [100]})
    try:
        if wait is not None and not wait.wait(30):
            raise TimeoutError('test process coordination')
        store.save(product, rates, inputs, {'expected_version': expected})
        queue.put('saved')
    except RuntimeError:
        queue.put('conflict')
    finally:
        if signal is not None:
            signal.set()


def test_two_processes_preserve_products_and_reject_stale_same_product(tmp_path, monkeypatch):
    import multiprocessing
    import __main__
    # AppTest can leave __main__ pointing at its generated app script. Windows
    # spawn must load this import-safe module instead of executing that script.
    monkeypatch.setattr(__main__, '__file__', __file__, raising=False)
    monkeypatch.setattr(__main__, '__spec__', None, raising=False)
    ctx = multiprocessing.get_context('spawn')
    path = tmp_path / 'concurrent.xlsx'
    queue = ctx.Queue()
    first = ctx.Process(target=_concurrent_publish, args=(path, 'A', None, None, None, queue))
    second = ctx.Process(target=_concurrent_publish, args=(path, 'B', None, None, None, queue))
    processes = [first, second]
    try:
        for process in processes:
            process.start()
        for process in processes:
            process.join(40)
            assert not process.is_alive() and process.exitcode == 0
        assert [queue.get(timeout=5), queue.get(timeout=5)] == ['saved', 'saved']
        store = PeriodTrendWorkbook(path)
        assert store.read('A') is not None and store.read('B') is not None
        version = store.read('A')[2]['version']
        signal = ctx.Event()
        first = ctx.Process(target=_concurrent_publish, args=(path, 'A', version, None, signal, queue))
        second = ctx.Process(target=_concurrent_publish, args=(path, 'A', version, signal, None, queue))
        processes = [first, second]
        for process in processes:
            process.start()
        for process in processes:
            process.join(40)
            assert not process.is_alive() and process.exitcode == 0
        assert sorted([queue.get(timeout=5), queue.get(timeout=5)]) == ['conflict', 'saved']
    finally:
        for process in processes:
            if process.is_alive():
                process.terminate()
                process.join(10)
        queue.close()
