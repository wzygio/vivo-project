"""Editable Inline weekly tables on the configured PowerPoint master."""
from __future__ import annotations

from copy import deepcopy
from datetime import timedelta
from io import BytesIO
from math import ceil
from typing import TYPE_CHECKING
from unicodedata import east_asian_width

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

if TYPE_CHECKING:
    from src.inline_domain.application.shared.weekly_summary import WeeklySummaryTable

BLUE = "0050FF"
ROW_COLORS = ("CDCEFF", "E8E8FF")
PRODUCTS_PER_PAGE = 8
BODY_TOP = 106
BODY_HEIGHT = 394


def _text(shape, text: str, size: float, *, color: str = "000000", bold: bool = False) -> None:
    shape.text = text
    for paragraph in shape.text_frame.paragraphs:
        paragraph.font.name = "Microsoft YaHei"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor.from_string(color)


def _cell(
    cell, text: str, *, header: bool = False, fill: str = "FFFFFF", font_size: float | None = None,
) -> None:
    cell.text = text
    cell.margin_left = cell.margin_right = Pt(5)
    cell.margin_top = cell.margin_bottom = Pt(3)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor.from_string(BLUE if header else fill)
    for p in cell.text_frame.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(font_size if font_size is not None else (15 if header else 13))
        p.font.bold = header
        p.font.color.rgb = RGBColor.from_string("FFFFFF" if header else "000000")
    cell.text_frame.word_wrap = True
    # python-pptx has no public table-cell border API.
    properties = cell._tc.get_or_add_tcPr()
    for side in ("L", "R", "T", "B"):
        line = OxmlElement(f"a:ln{side}")
        line.set("w", "10160")
        solid = OxmlElement("a:solidFill")
        rgb = OxmlElement("a:srgbClr")
        rgb.set("val", "FFFFFF")
        solid.append(rgb)
        line.append(solid)
        properties.append(line)


def _row_height(label: str, width: float) -> float:
    units = sum(1 if east_asian_width(c) in "WF" else .55 for c in label)
    height = max(35, ceil(units * 13 / (width - 10)) * 17 + 10)
    if height > BODY_HEIGHT - 38:
        raise ValueError("项目名称过长，无法放入报告表格。")
    return height


def _pages(rows, label_width: float):
    page, height = [], 38
    for row in rows:
        row_height = _row_height(row.label, label_width)
        if page and height + row_height > BODY_HEIGHT:
            yield page
            page, height = [], 38
        page.append((row, row_height))
        height += row_height
    if page or not rows:
        yield page


def _clone_template_slide(deck, source):
    """Copy slide-local assets while retaining the original layout/master."""
    slide = deck.slides.add_slide(source.slide_layout)
    for shape in list(slide.shapes):
        slide.shapes._spTree.remove(shape.element)
    relation_ids = {}
    for relationship in source.part.rels.values():
        if relationship.is_external:
            rid = slide.part.rels.get_or_add_ext_rel(relationship.reltype, relationship.target_ref)
        else:
            rid = slide.part.rels.get_or_add(relationship.reltype, relationship.target_part)
        relation_ids[relationship.rId] = rid
    for shape in source.shapes:
        element = deepcopy(shape.element)
        for node in element.iter():
            for attr, value in list(node.attrib.items()):
                if attr.startswith("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"):
                    node.set(attr, relation_ids[value])
        slide.shapes._spTree.insert_element_before(element, "p:extLst")
    background = source._element.cSld.find("{http://schemas.openxmlformats.org/presentationml/2006/main}bg")
    if background is not None:
        slide._element.cSld.insert(0, deepcopy(background))
    return slide


def _draw_table(slide, table, rows, product_start: int, width: float, category: str) -> None:
    array = table.factory == "ARRAY"
    products = table.products[product_start:product_start + PRODUCTS_PER_PAGE]
    compact = len(products) > 6
    widths = [72, 60, 210] if array else [72, 245]
    widths += [(width - sum(widths)) / len(products)] * len(products)
    headings = ["分类", "层别", "项目"] if array else ["分类", "项目"]
    headings += list(products)
    height = 38 + sum(h for _, h in rows)
    grid = slide.shapes.add_table(len(rows) + 1, len(headings), Pt(20), Pt(BODY_TOP),
                                  Pt(width), Pt(height)).table
    for i, column_width in enumerate(widths):
        grid.columns[i].width = Pt(column_width)
    grid.rows[0].height = Pt(38)
    for col, heading in enumerate(headings):
        _cell(grid.cell(0, col), heading, header=True, font_size=13 if compact else 15)
    for index, (row, row_height) in enumerate(rows, 1):
        grid.rows[index].height = Pt(row_height)
        values = [category, row.layer, row.label] if array else [category, row.label]
        values += list(row.values[product_start:product_start + len(products)])
        for col, value in enumerate(values):
            is_product = col >= (3 if array else 2)
            _cell(grid.cell(index, col), value, fill=ROW_COLORS[(index - 1) % 2],
                  font_size=12 if compact and is_product else 13)
    if len(rows) > 1:
        grid.cell(1, 0).merge(grid.cell(len(rows), 0))
        _cell(grid.cell(1, 0), category, fill=ROW_COLORS[1])
    if array:
        first = 1
        while first <= len(rows):
            last = first
            while last < len(rows) and rows[last][0].step_id == rows[first - 1][0].step_id:
                last += 1
            if last > first:
                grid.cell(first, 1).merge(grid.cell(last, 1))
                _cell(grid.cell(first, 1), rows[first - 1][0].layer, fill=ROW_COLORS[1])
            first = last + 1


def build_weekly_summary_ppt(
    tables: list[WeeklySummaryTable], *, template_bytes: bytes,
    category: str, metric_label: str, report_name: str,
) -> bytes:
    """Return a standalone PPTX. Neither template bytes nor business frames are changed."""
    if not tables or any(not table.products for table in tables):
        raise ValueError("请选择需要导出的产品和厂别。")
    deck = Presentation(BytesIO(template_bytes))
    if len(deck.slides) != 1:
        raise ValueError("周报模板需要包含一张基础页面。")
    # Keep the pristine source slide until every output slide has been cloned.
    source = deck.slides[0]
    width = deck.slide_width / 12700
    height = deck.slide_height / 12700
    if abs(width - 960) > 2 or abs(height - 540) > 2:
        raise ValueError("周报模板需要使用 16:9 的 960×540 页面。")
    page_number = 0
    for table in tables:
        for product_start in range(0, len(table.products), PRODUCTS_PER_PAGE):
            label_width = 210 if table.factory == "ARRAY" else 245
            for rows in _pages(table.rows, label_width):
                slide = _clone_template_slide(deck, source)
                page_number += 1
                title = slide.shapes.add_textbox(Pt(20), Pt(8), Pt(700), Pt(48))
                factory = "Array" if table.factory == "ARRAY" else table.factory
                _text(title, f"{factory} Inline监控Summary", 28, color=BLUE, bold=True)
                week = table.week_start.isocalendar()
                date_label = f"{week.year}-W{week.week:02d}  |  {table.week_start} ～ {table.week_end - timedelta(days=1)}"
                caption = slide.shapes.add_textbox(Pt(20), Pt(73), Pt(920), Pt(24))
                _text(caption, f"{date_label}    {metric_label}", 12, color="555555")
                if rows:
                    _draw_table(slide, table, rows, product_start, width - 40, category)
                else:
                    empty = slide.shapes.add_textbox(Pt(40), Pt(145), Pt(850), Pt(60))
                    products = "、".join(table.products[product_start:product_start + PRODUCTS_PER_PAGE])
                    _text(empty, f"{products}\n所选范围暂无 {category} 项目。", 18)
                footer = slide.shapes.add_textbox(Pt(20), Pt(514), Pt(920), Pt(20))
                _text(footer, f"{report_name} 周报    ·    {page_number}", 10, color="777777")
    # Remove only the unused, pristine template slide, retaining its master and assets.
    source_id = deck.slides._sldIdLst[0]
    deck.part.drop_rel(source_id.rId)
    deck.slides._sldIdLst.remove(source_id)
    deck.core_properties.title = f"{report_name} 周度报告"
    deck.core_properties.subject = metric_label
    deck.core_properties.author = ""
    deck.core_properties.last_modified_by = ""
    output = BytesIO()
    deck.save(output)
    return output.getvalue()
