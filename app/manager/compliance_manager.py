"""全指标矩阵的显示修饰；不修改计算结果，也不清理计算缓存。"""
from __future__ import annotations

import os
import tempfile
from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from src.shared_kernel.config import ConfigLoader
from src.shared_kernel.utils.excel_tools import read_workbook_sheet

CONFIG_PATH = ConfigLoader.get_compliance_config_path()
SHEET_NAME = "全指标预警看板"
ROW_COLUMN = "监控参数"


def parse_matrix_config(frame: pd.DataFrame) -> dict[tuple[str, str], bool]:
    """严格接受 True/False，禁止把非空字符串当作 True。"""
    if ROW_COLUMN not in frame or frame.columns.duplicated().any():
        raise ValueError("修饰配置应为监控参数 × 产品矩阵，不能使用旧四维规则表")
    frame = frame.dropna(how="all").copy()
    names = frame[ROW_COLUMN].astype("string").str.strip()
    if names.isna().any() or names.eq("").any() or names.duplicated().any():
        raise ValueError("修饰配置的监控参数不得为空或重复")
    switches = {}
    for index, name in names.items():
        for product in frame.columns:
            if product == ROW_COLUMN:
                continue
            value = str(frame.at[index, product]).strip().lower()
            if value not in ("true", "false"):
                raise ValueError(f"{name} × {product} 只能填写 True 或 False")
            switches[(str(name), str(product))] = value == "true"
    return switches


@st.cache_data(show_spinner=False, max_entries=16)
def _read_matrix_config(path: str, mtime_ns: int, size: int) -> dict:
    return parse_matrix_config(read_workbook_sheet(Path(path), SHEET_NAME))


def load_matrix_config(path: Path | None = None) -> dict[tuple[str, str], bool]:
    source = Path(path or CONFIG_PATH)
    stat = source.stat()
    return _read_matrix_config(str(source.resolve()), stat.st_mtime_ns, stat.st_size)


def apply_matrix_compliance(payload: Mapping[str, Any], switches: Mapping) -> dict:
    """在缓存外复制单元格；表格与详情必须消费同一显示 payload。"""
    names = {row["row_key"]: row["display_name"] for row in payload.get("rows", [])}
    cells = {}
    for (row_key, product), cell in payload.get("cells", {}).items():
        if switches.get((names.get(row_key), product), False):
            cells[(row_key, product)] = {**cell, "state": "ok", "message": ""}
        else:
            cells[(row_key, product)] = dict(cell)
    return {**payload, "cells": cells}


def write_matrix_template(
    rows: Sequence[Mapping[str, Any]], products: Sequence[str], path: Path | None = None,
) -> None:
    """显式初始化/替换模板；查询和普通刷新绝不重写用户开关。"""
    target = Path(path or CONFIG_PATH)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append([ROW_COLUMN, *products])
    for row in rows:
        sheet.append([row["display_name"], *([False] * len(products))])
    sheet.freeze_panes = "B2"
    sheet.column_dimensions["A"].width = 34
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    for column in range(2, len(products) + 2):
        letter = sheet.cell(1, column).column_letter
        sheet.column_dimensions[letter].width = 15
        for row_index in range(2, sheet.max_row + 1):
            cell = sheet.cell(row_index, column)
            cell.alignment = Alignment(horizontal="center")
            cell.fill = PatternFill("solid", fgColor="F0F5FA" if row_index % 2 == 0 else "FFFFFF")
    validation = DataValidation(type="list", formula1='"True,False"', allow_blank=False)
    validation.showErrorMessage = True
    validation.error = "请填写 True 或 False"
    sheet.add_data_validation(validation)
    if rows and products:
        validation.add(f"B2:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}")
    descriptor, temporary = tempfile.mkstemp(suffix=".xlsx", dir=target.parent)
    os.close(descriptor)
    try:
        workbook.save(temporary)
        os.replace(temporary, target)
    finally:
        workbook.close()
        Path(temporary).unlink(missing_ok=True)


def get_compliance_file_signature() -> str:
    """旧明细缓存兼容接口：显示修饰文件不得进入计算缓存签名。"""
    return "display-only-compliance"
